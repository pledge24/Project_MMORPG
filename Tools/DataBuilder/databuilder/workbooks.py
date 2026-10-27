"""워크북을 읽어 표와 설정으로 만들고, 시트 규칙과 표끼리의 규칙을 검사한다.

가로형 표의 시트 규칙: 1행 이름, 2행 타입, 3행 대상, 4행 설명, 5행부터 데이터.
한 워크북의 시트는 1~3행이 같아야 하고, 데이터 행을 시트 순서대로 이어 붙인다. 시트 이름은 데이터가 아니다.

세로형 설정 시트: 1행이 `Key | Type | Target | Value | Description`이고, 2행부터 한 행이 필드 하나다.

오류는 모두 모아 돌려준다. 오류 하나는 워크북, 시트, 행, 열을 적은 한 줄이다.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

from databuilder.column_types import CellError, ColumnType, convert_cell, is_blank, parse_type
from databuilder.config import Config, SettingConfig, TableConfig
from databuilder.enums import EnumDef

COLUMN_NAME = re.compile(r"^[A-Z][A-Za-z0-9]*$")
TARGETS = ("both", "client", "server", "none")
SETTING_HEADER = ("Key", "Type", "Target", "Value", "Description")
FIRST_DATA_ROW = 5


def targets_server(target: str) -> bool:
    return target in ("both", "server")


def targets_client(target: str) -> bool:
    return target in ("both", "client")


@dataclass
class Column:
    name: str
    type: ColumnType
    target: str
    description: str
    letter: str


@dataclass
class Row:
    values: dict[str, Any]
    sheet: str
    number: int


@dataclass
class Table:
    config: TableConfig
    columns: list[Column]
    rows: list[Row] = field(default_factory=list)

    @property
    def name(self) -> str:
        return self.config.name

    @property
    def key_columns(self) -> list[Column]:
        return [column for column in self.columns if column.type.is_key]


@dataclass
class Setting:
    config: SettingConfig
    fields: list[Column]
    values: dict[str, Any] = field(default_factory=dict)
    # 필드 이름 → 시트의 행 번호. 오류 위치에 쓴다.
    rows: dict[str, int] = field(default_factory=dict)

    @property
    def name(self) -> str:
        return self.config.name


@dataclass
class DesignData:
    """검증을 마친 원본 전체."""

    tables: dict[str, Table]
    settings: dict[str, Setting]
    enums: dict[str, EnumDef]


def locate(workbook: str, sheet: str | None = None, row: int | None = None, letter: str | None = None, name: str | None = None) -> str:
    """오류 위치. `Currency.xlsx 시트 '재화' 5행 C열(Type)`."""
    parts = [workbook]
    if sheet is not None:
        parts.append(f"시트 '{sheet}'")
    if row is not None:
        parts.append(f"{row}행")
    if letter is not None:
        parts.append(f"{letter}열" + (f"({name})" if name else ""))
    return " ".join(parts)


def cell_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def load_workbook(path: Path, errors: list[str]):
    if not path.is_file():
        errors.append(f"{path.name}: 워크북이 없다 ({path})")
        return None
    try:
        return openpyxl.load_workbook(path, data_only=True)
    except Exception as error:  # openpyxl은 잘못된 파일에 여러 종류의 예외를 던진다.
        errors.append(f"{path.name}: 워크북을 열지 못했다: {error}")
        return None


def read_table_layout(config: TableConfig, source_dir: Path, errors: list[str]) -> tuple[Table | None, list[tuple[str, int, tuple]]]:
    """시트 1~4행으로 열을 정하고, 데이터 행의 셀을 그대로 모은다. 값 변환은 열 타입을 모두 안 뒤에 한다."""
    workbook = load_workbook(source_dir / config.workbook, errors)
    if workbook is None:
        return None, []

    columns: list[Column] | None = None
    first_sheet = ""
    raw_rows: list[tuple[str, int, tuple]] = []
    error_count = len(errors)

    for sheet in workbook.worksheets:
        header = [list(row) for row in sheet.iter_rows(min_row=1, max_row=4, values_only=True)]
        header += [[] for _ in range(4 - len(header))]
        names = header[0]
        width = 0
        while width < len(names) and not is_blank(names[width]):
            width += 1
        if width == 0:
            errors.append(f"{locate(config.workbook, sheet.title, 1)}: 1행에 열 이름이 없다")
            continue

        def at(row: list, index: int) -> Any:
            return row[index] if index < len(row) else None

        sheet_columns: list[Column] = []
        for index in range(width):
            letter = get_column_letter(index + 1)
            name = cell_text(names[index])
            where_name = locate(config.workbook, sheet.title, 1, letter, name)
            if not COLUMN_NAME.match(name):
                errors.append(f"{where_name}: 열 이름은 PascalCase 식별자로 쓴다")
            if any(column.name == name for column in sheet_columns):
                errors.append(f"{where_name}: 열 이름이 겹친다")
            try:
                column_type = parse_type(at(header[1], index))
            except CellError as error:
                errors.append(f"{locate(config.workbook, sheet.title, 2, letter, name)}: {error}")
                column_type = ColumnType("string")
            target = cell_text(at(header[2], index))
            if target not in TARGETS:
                errors.append(f"{locate(config.workbook, sheet.title, 3, letter, name)}: 대상 '{target}'은 {', '.join(TARGETS)} 중 하나여야 한다")
            sheet_columns.append(Column(name, column_type, target, cell_text(at(header[3], index)), letter))

        if columns is None:
            columns, first_sheet = sheet_columns, sheet.title
        elif [(c.name, c.type, c.target) for c in sheet_columns] != [(c.name, c.type, c.target) for c in columns]:
            errors.append(f"{locate(config.workbook, sheet.title)}: 열 구성(1~3행)이 시트 '{first_sheet}'와 다르다")

        for number, row in enumerate(sheet.iter_rows(min_row=FIRST_DATA_ROW, values_only=True), start=FIRST_DATA_ROW):
            if all(is_blank(value) for value in row):
                continue
            for index in range(width, len(row)):
                if not is_blank(row[index]):
                    errors.append(f"{locate(config.workbook, sheet.title, number, get_column_letter(index + 1))}: 이름 없는 열에 값이 있다")
            raw_rows.append((sheet.title, number, tuple(row[:width]) + (None,) * (width - len(row))))

    if columns is None:
        if len(errors) == error_count:
            errors.append(f"{config.workbook}: 시트가 없다")
        return None, []

    table = Table(config, columns)
    check_table_columns(table, errors)
    return table, raw_rows


def check_table_columns(table: Table, errors: list[str]) -> None:
    workbook = table.config.workbook
    keys = table.key_columns
    if not keys:
        errors.append(f"{workbook}: 키 열이 없다. 키 열의 타입은 'key:'로 시작한다")

    # 키는 행을 가리키는 값이라, 열을 하나라도 받는 쪽에는 키도 가야 한다.
    for side, accepts in (("server", targets_server), ("client", targets_client)):
        if any(accepts(column.target) for column in table.columns):
            for key in keys:
                if not accepts(key.target):
                    errors.append(f"{locate(workbook, None, 3, key.letter, key.name)}: {side}가 받는 열이 있으므로 키 열도 {side}에 가야 한다")

    parents = [column for column in table.columns if column.type.kind == "parent"]
    if len(parents) > 1:
        errors.append(f"{workbook}: 부모 표 참조 열은 하나뿐이어야 한다({', '.join(c.name for c in parents)})")
    for column in parents:
        if column.type.target == table.name:
            errors.append(f"{locate(workbook, None, 2, column.letter, column.name)}: 자기 자신을 부모 표로 삼을 수 없다")


def read_setting_layout(config: SettingConfig, source_dir: Path, errors: list[str]) -> tuple[Setting | None, list[tuple[Column, int, Any]]]:
    workbook = load_workbook(source_dir / config.workbook, errors)
    if workbook is None:
        return None, []
    if config.sheet not in workbook.sheetnames:
        errors.append(f"{config.workbook}: 시트 '{config.sheet}'가 없다")
        return None, []

    sheet = workbook[config.sheet]
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    header = tuple(cell_text(value) for value in (rows[0] if rows else [])[: len(SETTING_HEADER)])
    if header != SETTING_HEADER:
        errors.append(f"{locate(config.workbook, config.sheet, 1)}: 세로형 설정 시트의 1행은 {' | '.join(SETTING_HEADER)}다")
        return None, []

    setting = Setting(config, [])
    raw_values: list[tuple[Column, int, Any]] = []
    for number, row in enumerate(rows[1:], start=2):
        row = row + [None] * (len(SETTING_HEADER) - len(row))
        if all(is_blank(value) for value in row):
            continue
        name = cell_text(row[0])
        if not COLUMN_NAME.match(name):
            errors.append(f"{locate(config.workbook, config.sheet, number, 'A', name)}: 키는 PascalCase 식별자로 쓴다")
        if name in setting.rows:
            errors.append(f"{locate(config.workbook, config.sheet, number, 'A', name)}: 키가 {setting.rows[name]}행과 겹친다")
        try:
            column_type = parse_type(row[1])
            if column_type.is_key or column_type.kind == "parent":
                raise CellError(f"설정 시트에는 키와 부모 표 참조를 쓰지 않는다: '{row[1]}'")
        except CellError as error:
            errors.append(f"{locate(config.workbook, config.sheet, number, 'B', name)}: {error}")
            column_type = ColumnType("string")
        target = cell_text(row[2])
        if target not in TARGETS:
            errors.append(f"{locate(config.workbook, config.sheet, number, 'C', name)}: 대상 '{target}'은 {', '.join(TARGETS)} 중 하나여야 한다")
        column = Column(name, column_type, target, cell_text(row[4]), "D")
        setting.fields.append(column)
        setting.rows[name] = number
        raw_values.append((column, number, row[3]))
    return setting, raw_values


def resolve_column_type(column: Column, where: str, data: DesignData, errors: list[str]) -> ColumnType | None | bool:
    """열 타입이 가리키는 열거형과 표를 확인한다.

    참조 열이면 값의 타입(참조를 따라가 닿은 키의 정수, 문자열, 열거형 타입)을, 아니면 None을, 틀리면 False를 돌려준다.
    """
    column_type = column.type
    visited: list[str] = []
    while True:
        if column_type.kind == "enum" and column_type.target not in data.enums:
            errors.append(f"{where}: 열거형 '{column_type.target}'이 .proto에도 설정 파일에도 없다")
            return False
        if not column_type.is_reference:
            return None if column_type is column.type else ColumnType(column_type.kind, column_type.target)
        target = data.tables.get(column_type.target)
        if target is None:
            errors.append(f"{where}: 참조하는 표 '{column_type.target}'이 설정 파일에 없다")
            return False
        if len(target.key_columns) != 1:
            errors.append(f"{where}: 표 '{target.name}'은 키 열이 하나가 아니라서 참조할 수 없다")
            return False
        if target.name in visited:
            errors.append(f"{where}: 키 참조가 표 '{target.name}'으로 되돌아온다")
            return False
        visited.append(target.name)
        column_type = target.key_columns[0].type


def value_type(column_type: ColumnType, data: DesignData) -> ColumnType:
    """참조를 따라가 닿은 키의 타입. 참조가 아니면 그대로다. 검사를 통과한 원본에만 쓴다."""
    while column_type.is_reference:
        column_type = data.tables[column_type.target].key_columns[0].type
    return column_type


def sort_key(table: Table, row: Row, data: DesignData) -> tuple:
    """키 순 정렬에 쓴다. 열거형 키(열거형 키를 가진 표의 참조 포함)는 값 번호 순서다."""
    parts = []
    for column in table.key_columns:
        value = row.values[column.name]
        key_type = value_type(column.type, data)
        if key_type.kind == "enum":
            value = data.enums[key_type.target].values[value]
        parts.append(value)
    return tuple(parts)


def read_design_data(config: Config, enums: dict[str, EnumDef], errors: list[str]) -> DesignData:
    """설정 파일이 가리키는 워크북을 모두 읽고 검사한다. 오류는 errors에 더한다."""
    data = DesignData({}, {}, enums)
    raw_tables: dict[str, list[tuple[str, int, tuple]]] = {}
    for table_config in config.tables.values():
        table, raw_rows = read_table_layout(table_config, config.source_dir, errors)
        if table is not None:
            data.tables[table.name] = table
            raw_tables[table.name] = raw_rows
    raw_settings: dict[str, list[tuple[Column, int, Any]]] = {}
    for setting_config in config.settings.values():
        setting, raw_values = read_setting_layout(setting_config, config.source_dir, errors)
        if setting is not None:
            data.settings[setting.name] = setting
            raw_settings[setting.name] = raw_values

    # 열 타입이 모두 정해진 뒤에야 참조 열의 값 타입(참조하는 표의 키 타입)을 안다.
    for table in data.tables.values():
        resolved = {}
        for column in table.columns:
            resolved[column.name] = resolve_column_type(column, locate(table.config.workbook, None, 2, column.letter, column.name), data, errors)
        for sheet, number, cells in raw_tables[table.name]:
            row = Row({}, sheet, number)
            for column, value in zip(table.columns, cells):
                if resolved[column.name] is False:
                    continue
                if column.target == "none" and is_blank(value):
                    row.values[column.name] = None  # 어디에도 싣지 않는 열(메모)은 비워 둘 수 있다.
                    continue
                try:
                    row.values[column.name] = convert_cell(value, column.type, enums, resolved[column.name] or None)
                except CellError as error:
                    errors.append(f"{locate(table.config.workbook, sheet, number, column.letter, column.name)}: {error}")
            table.rows.append(row)

    for setting in data.settings.values():
        for column, number, value in raw_settings[setting.name]:
            where = locate(setting.config.workbook, setting.config.sheet, number, "D", column.name)
            resolved = resolve_column_type(column, where, data, errors)
            if resolved is False:
                continue
            if column.target == "none" and is_blank(value):
                setting.values[column.name] = None
                continue
            try:
                setting.values[column.name] = convert_cell(value, column.type, enums, resolved or None)
            except CellError as error:
                errors.append(f"{where}: {error}")

    check_keys(data, errors)
    check_references(data, errors)
    return data


def check_keys(data: DesignData, errors: list[str]) -> None:
    for table in data.tables.values():
        keys = table.key_columns
        if not keys:
            continue
        seen: dict[tuple, Row] = {}
        for row in table.rows:
            if any(column.name not in row.values for column in keys):
                continue  # 키 셀 자체가 틀렸고 그 오류는 이미 있다.
            key = tuple(row.values[column.name] for column in keys)
            if key in seen:
                first = seen[key]
                names = ", ".join(column.name for column in keys)
                shown = ", ".join(str(value) for value in key)
                errors.append(
                    f"{locate(table.config.workbook, row.sheet, row.number)}: 키({names}) {shown}이 "
                    f"시트 '{first.sheet}' {first.number}행과 겹친다"
                )
            else:
                seen[key] = row


def table_keys(table: Table) -> set:
    key = table.key_columns[0].name
    return {row.values[key] for row in table.rows if key in row.values}


def check_references(data: DesignData, errors: list[str]) -> None:
    def check(values: Any, column: Column, where: str) -> None:
        target = data.tables.get(column.type.target)
        if target is None or len(target.key_columns) != 1:
            return  # 열 타입 오류로 이미 걸렸다.
        known = table_keys(target)
        for value in values if isinstance(values, list) else [values]:
            if value is not None and value not in known:
                errors.append(f"{where}: {value}가 표 '{target.name}'의 키({target.key_columns[0].name})에 없다")

    for table in data.tables.values():
        for column in table.columns:
            if not column.type.is_reference:
                continue
            for row in table.rows:
                if column.name in row.values:
                    check(row.values[column.name], column, locate(table.config.workbook, row.sheet, row.number, column.letter, column.name))
    for setting in data.settings.values():
        for column in setting.fields:
            if column.type.is_reference and column.name in setting.values:
                where = locate(setting.config.workbook, setting.config.sheet, setting.rows[column.name], "D", column.name)
                check(setting.values[column.name], column, where)
