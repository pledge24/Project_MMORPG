"""테스트용 원본을 임시 폴더에 만든다. 워크북, `.proto`, 설정 파일을 테스트마다 새로 쓴다."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from typing import Any

import openpyxl

from databuilder.build import BuildResult, generate
from databuilder.config import load_config

PROTO = """// @packet_id_range 100 199
syntax = "proto3";
package Protocol;

// 아이템 종류
enum ItemType
{
    ITEM_TYPE_NONE = 0;
    ITEM_TYPE_GEAR = 1;
    ITEM_TYPE_CONSUMABLE = 2;
}

enum CharacterClass
{
    CHARACTER_CLASS_NONE = 0;
    CHARACTER_CLASS_WARRIOR = 1;
    CHARACTER_CLASS_MAGE = 2;
}

message S_TEST
{
    enum Inner
    {
        NOT_READ = 0;
    }
}
"""


def sheet(names: list[str], types: list[str], targets: list[str], *rows: list[Any], descriptions: list[str] | None = None) -> list[list[Any]]:
    """가로형 시트의 행 목록. 4행 설명은 비워 두면 빈 칸이다."""
    return [names, types, targets, descriptions or [None] * len(names), *rows]


def setting_sheet(*fields: list[Any]) -> list[list[Any]]:
    """세로형 설정 시트의 행 목록. 필드는 [Key, Type, Target, Value, Description]이다."""
    return [["Key", "Type", "Target", "Value", "Description"], *fields]


class DesignDataCase(unittest.TestCase):
    """임시 저장소 하나를 두고 워크북과 설정을 써서 DataBuilder를 돌린다."""

    def setUp(self) -> None:
        self._temp = tempfile.TemporaryDirectory(prefix="databuilder-test-")
        self.root = Path(self._temp.name)
        (self.root / "DesignData").mkdir()
        (self.root / "Schema").mkdir()
        (self.root / "Schema" / "Common.proto").write_bytes(PROTO.encode("cp949"))
        self.out_dir = self.root / "Server" / "Game" / "Generated"
        self.tables: dict[str, tuple[str, list[str]]] = {}
        self.settings: dict[str, tuple[str, str, list[str]]] = {}
        self.enums: dict[str, dict[str, int]] = {"CurrencyType": {"GOLD": 1, "HONOR": 2}}

    def tearDown(self) -> None:
        self._temp.cleanup()

    def write_workbook(self, name: str, sheets: dict[str, list[list[Any]]]) -> Path:
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        for title, rows in sheets.items():
            worksheet = workbook.create_sheet(title)
            for row in rows:
                worksheet.append(row)
        path = self.root / "DesignData" / name
        workbook.save(path)
        return path

    def add_table(self, name: str, *sheets: list[list[Any]], modules: list[str] | None = None) -> None:
        """표 하나를 워크북 `<name>.xlsx`로 쓰고 설정에 넣는다. 시트 이름은 Sheet1, Sheet2...다."""
        self.write_workbook(f"{name}.xlsx", {f"Sheet{i + 1}": rows for i, rows in enumerate(sheets)})
        self.tables[name] = (f"{name}.xlsx", modules or ["game"])

    def add_setting(self, name: str, rows: list[list[Any]], workbook: str = "Config.xlsx") -> None:
        path = self.root / "DesignData" / workbook
        existing = {}
        if path.exists():
            loaded = openpyxl.load_workbook(path)
            existing = {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in loaded.worksheets}
        existing[name] = rows
        self.write_workbook(workbook, existing)
        self.settings[name] = (workbook, name, ["game"])

    def config_path(self) -> Path:
        lines = ['source_dir = "DesignData"', 'proto_dir = "Schema"', "", "[modules.game]", 'server_dir = "Server/Game/Generated"', ""]
        for name, (workbook, modules) in self.tables.items():
            lines += [f"[tables.{name}]", f'workbook = "{workbook}"', f"modules = {modules!r}".replace("'", '"'), ""]
        for name, (workbook, sheet_name, modules) in self.settings.items():
            lines += [f"[settings.{name}]", f'workbook = "{workbook}"', f'sheet = "{sheet_name}"', f"modules = {modules!r}".replace("'", '"'), ""]
        for name, values in self.enums.items():
            lines += [f"[enums.{name}]"] + [f"{value} = {number}" for value, number in values.items()] + [""]
        path = self.root / "databuilder.toml"
        path.write_text("\n".join(lines), encoding="utf-8")
        return path

    def build(self) -> BuildResult:
        return generate(load_config(self.config_path(), self.root))

    def build_ok(self) -> BuildResult:
        result = self.build()
        self.assertEqual(result.errors, [], "통과해야 하는 원본에서 오류가 났다")
        return result

    def output(self, file_name: str) -> str:
        return self.build_ok().outputs[self.out_dir][file_name]

    def assert_error(self, *fragments: str) -> list[str]:
        """오류 중 하나가 조각을 모두 담는지 본다."""
        errors = self.build().errors
        self.assertTrue(errors, "실패해야 하는 원본이 통과했다")
        matching = [error for error in errors if all(fragment in error for fragment in fragments)]
        self.assertTrue(matching, f"{fragments}를 모두 담은 오류가 없다:\n" + "\n".join(errors))
        return errors
